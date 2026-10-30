from pathlib import Path
import openpyxl,json,hashlib,datetime,collections,decimal,re
import argparse
p=argparse.ArgumentParser(description='Read-only, local private preview; never connects to an ERP or writes the source workbook.')
p.add_argument('--source',required=True,type=Path);p.add_argument('--output',required=True,type=Path);a=p.parse_args()
P=a.source;D=a.output;D.mkdir(exist_ok=False)
sha=hashlib.sha256(P.read_bytes()).hexdigest();w=openpyxl.load_workbook(P,read_only=True,data_only=False);v=openpyxl.load_workbook(P,read_only=True,data_only=True)
def serial(x):
 if isinstance(x,(datetime.date,datetime.datetime,datetime.time)):return x.isoformat()
 if isinstance(x,decimal.Decimal):return str(x)
 if isinstance(x,openpyxl.worksheet.formula.ArrayFormula):return {"formula":x.text,"range":x.ref,"array":True}
 return x
mapping={
 'A':('operation.reference_code','group_candidate; verify unique entity before import'),'B':('ar.invoice_reference','invoice candidate, not CFDI UUID'),'C':('ar.issue_date','date validation'), 'D':('ar.credit_days','nonnegative integer'), 'E':('ar.payment_date_evidence','not sufficient to create payment'),
 'F':('ar.line.service','numeric service charge'), 'G':('ar.line.overweight','numeric extra charge'),'H':('ar.line.crossing','numeric extra charge'),'I':('ar.line.other','numeric extra charge'), 'J':('ar.fx_rate','direction/date pending'), 'K':('ar.subtotal_control','formula and cache retained'), 'L':('ar.tax_control','rate/base validation'),'M':('ar.total_control','by currency only'),'N':('ar.currency','explicit MXN/USD; never infer'),
 'O':('ap.provider_candidate','resolve actual provider ID'), 'P':('ap.invoice_reference','group per provider+invoice+currency; validate'),'Q':('ap.issue_date','date validation'),'R':('ap.line.service','numeric'),'S':('ap.line.overweight','numeric'),'T':('ap.line.crossing','numeric'),'U':('ap.line.other','numeric'),'V':('ap.subtotal_control','formula/cache'),'W':('ap.withholding_control','base/sign validation'),'X':('ap.tax_control','base/rate validation'),'Y':('ap.total_control','by currency only'),'Z':('ap.currency','explicit only'),'AA':('ap.fx_rate','direction/date pending'),'AB':('ap.due_date','formula/date validation'),
 'AC':('operation.route_evidence','not structured origin/destination'),'AD':('customer_candidate','ID reconciliation'), 'AE':('operation.goods_evidence','not financial notes'),'AF':('operation.material_evidence','catalog pending'), 'AG':('operation.milestone_evidence','multiple dates require split/meaning'), 'AH':('operation.equipment_evidence','third_party optional'), 'AI':('operation.unit_evidence','not unique operation ID'),'AJ':('operation.measurement_evidence','parse unit/value explicitly'),'AK':('operation.customs_evidence','retain source identifier'),'AL':('operation.external_reference','not guaranteed unique'),'AM':('ar.status_evidence','not automatic operational status'), 'AN':('ap.payment_date_evidence','does not imply amount/application'),'AO':('restricted.commission_control','no import rule'), 'AP':('restricted.profit_control','reconcile AR/AP FX; not independent amount'),'AQ':('restricted.margin_control','denominator/FX rule pending'),'AR':('restricted.payroll_label','no payroll import authorized'),'AS':('restricted.payroll_control','independent schedule; not one amount per operation'),'AT':('restricted.payroll_allocation_control','allocation unresolved'),'AU':('restricted.net_profit_control','formula only; not factual ledger posting')}
summary={'source_sha256':sha,'read_only':True,'imports_executed':0,'recalculated':False,'sheets':[],'mapping':{k:{'target':a,'rule':b} for k,(a,b) in mapping.items()}}
for s in w:
 if s.title not in ('ROTERO','ROTERITO'):continue
 cached=v[s.title];stats=collections.Counter();types=collections.defaultdict(collections.Counter);totals=collections.defaultdict(decimal.Decimal);refs=collections.defaultdict(list);fingerprints=collections.defaultdict(list);last=None
 company={'sheet':s.title,'rfc':'WLS110325JH8' if s.title=='ROTERO' else 'WRO2507109E0','legal_name':"WORLD'S LOGISTIC SOLUTIONS ROTERO" if s.title=='ROTERO' else 'WLS ROTERO'}
 with (D/(s.title+'-preview.jsonl')).open('w',encoding='utf8') as f:
  for row,crow in zip(s.iter_rows(min_row=2,max_col=s.max_column),cached.iter_rows(min_row=2,max_col=s.max_column)):
   n=row[0].row if hasattr(row[0],'row') else stats['last_row']+1;stats['last_row']=n
   cells={openpyxl.utils.get_column_letter(i):{'cell':c.coordinate,'type':c.data_type,'value':serial(c.value),'cached':serial(crow[i-1].value),'number_format':c.number_format,'cached_type':crow[i-1].data_type} for i,c in enumerate(row,1) if c.value is not None}
   if not cells:stats['blank_rows']+=1;continue
   form=[c for c in cells.values() if c['type']=='f'];literal=[c for c in cells.values() if c['type']!='f'];flags=[]
   stats['formula_cells']+=len(form);stats['formula_cache_missing']+=sum(c['cached'] is None for c in form);stats['excel_errors']+=sum(c['type']=='e' or c['cached_type']=='e' for c in cells.values())
   for k,c in cells.items():types[k][c['type']]+=1
   folio=str(cells.get('A',{}).get('value') or '').strip()
   if folio=='Folio WLS':kind='repeated_header'
   elif folio:kind='anchor_candidate';last=n;refs[folio].append(n)
   elif literal and any(k in cells for k in ['B','O','P','AD','AG','AI','AK','AL']):kind='continuation_candidate' if last else 'unanchored_candidate';flags+=['grouping_requires_review']
   else:kind='formula_or_auxiliary_schedule';flags+=['not_operation']
   stats[kind]+=1
   if kind in ('anchor_candidate','continuation_candidate','unanchored_candidate'):
    for amount,currency,leg in [('M','N','ar'),('Y','Z','ap')]:
     val=cells.get(amount,{}).get('cached');cur=str(cells.get(currency,{}).get('cached') or '').strip().upper()
     if isinstance(val,(int,float)) and not isinstance(val,bool):totals[leg+':'+(cur if cur in ('MXN','USD') else 'UNKNOWN')]+=decimal.Decimal(str(val))
     elif val is not None:flags.append(leg+'_nonnumeric_total')
     if cur and cur not in ('MXN','USD'):flags.append(leg+'_currency_unresolved')
    for col in ['C','E','Q','AB','AN']:
     x=cells.get(col,{}).get('cached')
     if x is not None and not isinstance(crow[openpyxl.utils.column_index_from_string(col)-1].value,(datetime.datetime,datetime.date)):flags.append('date_review_'+col)
    if isinstance(cells.get('AG',{}).get('value'),str):flags.append('milestones_text_requires_review')
    if not folio:flags.append('no_explicit_folio')
   h=hashlib.sha256(json.dumps({k:{a:c[a] for a in ['type','value','cached']} for k,c in cells.items()},sort_keys=True,ensure_ascii=False).encode()).hexdigest();fingerprints[h].append(n)
   record={'source_sha256':sha,'source_sheet':s.title,'source_row':n,'company_proposal':company,'row_fingerprint':h,'idempotency_key':hashlib.sha256((sha+'|'+s.title+'|'+str(n)+'|'+h).encode()).hexdigest(),'kind':kind,'previous_anchor_candidate':last if not folio else None,'flags':flags,'cells':cells,'status':'HOLD; no approved entity grouping, IDs or ledger reconciliation'}
   f.write(json.dumps(record,ensure_ascii=False)+'\n')
 duplicates=[{'rows':rr,'kind':'same_nonblank_folio_not_auto_merge'} for rr in refs.values() if len(rr)>1];exact=[rr for rr in fingerprints.values() if len(rr)>1]
 summary['sheets'].append({'company_proposal':company,'dimensions':[s.max_row,s.max_column],'counts':dict(stats),'column_types':{k:dict(t) for k,t in types.items()},'duplicate_folio_row_groups':duplicates,'identical_row_groups':exact,'raw_candidate_totals_by_leg_currency':{k:str(x) for k,x in totals.items()},'totals_status':'unreconciled source controls; possible invoice repetitions/continuations; not posting totals'})
(D/'manifest.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
(D/'LEEME.md').write_text('''# Previsualización privada, sin importación

Las filas JSONL conservan hoja/fila/celda, tipo, fórmula original y caché. No se recalculó ni alteró el Excel. Las hojas proponen empresa por procedencia aprobada, sin crear IDs ni empresas. Todos los registros están en HOLD.

Antes de un lote: aprobar agrupación Folio/continuaciones; conciliar facturas repetidas por emisor/proveedor, referencia y moneda; resolver IDs de cliente/proveedor; desglosar fechas múltiples; verificar bases de IVA/retención, dirección/tipo de cambio y asignación de nómina. No crear pagos a partir de una fecha sin importe/aplicación. AP/AQ/AU son controles derivados, no nuevos asientos. AR:AU pertenecen además a un calendario auxiliar que no debe multiplicarse por operaciones.

Totales en manifest: sumas crudas de M/Y sólo en candidatos, separadas por moneda. No son totales de un lote aprobado. Separar las filas repetidas y controles antes de conciliarlos contra cada documento. Comparar F:I→K, K+L→M y R:U→V, V-W+X→Y usando valores/monedas explícitos. Comparar utilidad/margen sólo después de moneda y nómina. Caches pueden ser obsoletas o erróneas; revisar en Excel sin alterar fuente.

Idempotencia propuesta: huella de archivo+hoja+fila+contenido impide repetir exactamente esta instantánea. Una versión modificada requiere una tabla de correspondencia aprobada de entidades y revisión de conflictos, no una nueva alta automática. La huella no sustituye una clave de negocio. El manifiesto exacto y las decisiones se aprobarán antes de cualquier importación, con rollback del lote por IDs y conservación de trazabilidad. Ningún dato real se envió al repositorio ni a Supabase.
''',encoding='utf8')
assert hashlib.sha256(P.read_bytes()).hexdigest()==sha
print(json.dumps({'source_unchanged':True,'sheets':[{'sheet':x['company_proposal']['sheet'],'counts':x['counts'],'duplicate_folio_groups':len(x['duplicate_folio_row_groups']),'identical_row_groups':len(x['identical_row_groups'])} for x in summary['sheets']],'private_output':str(D)},ensure_ascii=True))
